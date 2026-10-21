"""Make the three-page editable student explainer from genuine model evidence."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'Final_Delivery'
ASSET=OUT/'Presentation_Assets'
SCOPE=OUT/'Results'/'05_Scope_Captures'
DEST=OUT/'Vataleni_PMSG_SSO_Three_Page_Explainer.docx'
IMG_DIR=ROOT/'tmp'/'three_page_images'
IMG_DIR.mkdir(parents=True,exist_ok=True)

def sized(path,max_width=3000):
    image=Image.open(path)
    if image.width<=max_width:
        return path
    target=IMG_DIR/path.name
    height=round(image.height*max_width/image.width)
    image.resize((max_width,height),Image.Resampling.LANCZOS).save(target,optimize=True)
    return target

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(11.69);sec.page_height=Inches(8.27)
sec.top_margin=Inches(.55);sec.bottom_margin=Inches(.51)
sec.left_margin=Inches(.65);sec.right_margin=Inches(.65)
sec.header_distance=Inches(.22);sec.footer_distance=Inches(.22)
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10)
normal.font.color.rgb=RGBColor(37,50,62)
normal.paragraph_format.space_after=Pt(5)
for name, size in [('Title',18),('Heading 1',15),('Heading 2',11)]:
    st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size)
    st.font.bold=True;st.font.color.rgb=RGBColor(19,43,67)
    st.paragraph_format.space_after=Pt(6)
caption=doc.styles.add_style('Figure Caption',WD_STYLE_TYPE.PARAGRAPH)
caption.font.name='Arial';caption.font.size=Pt(8.5);caption.font.italic=True
caption.font.color.rgb=RGBColor(70,82,94)
caption.paragraph_format.space_after=Pt(5)

footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.style=doc.styles['Normal']
run=footer.add_run('Vataleni PMSG–SSO  |  ');run.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def title(s,newpage=False):
    p=doc.add_paragraph(style='Title')
    p.paragraph_format.page_break_before=newpage
    p.add_run(s)
def body(s):
    doc.add_paragraph(s)
def cap(s):
    p=doc.add_paragraph(style='Figure Caption');p.add_run(s)
def picture(path,width):
    assert path.is_file(),path
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(2)
    p.add_run().add_picture(str(sized(path)),width=Inches(width))
def two_images(paths,widths):
    table=doc.add_table(rows=1,cols=2);table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    for i,(path,width) in enumerate(zip(paths,widths)):
        cell=table.cell(0,i);cell.width=Inches(5.1)
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].add_run().add_picture(str(sized(path,1800)),width=Inches(width))

# Page 1 - architecture.
title('1  Physical model architecture')
body('The saved Simulink model connects a wind turbine and single-mass shaft to a native PMSG, machine-side converter (MSC), DC capacitor, grid-side converter (GSC), RL filter, point of common coupling (PCC) and three-phase grid. The components in Figure 1 are the executable Simscape network.')
picture(ASSET/'01_complete_system.png',10.05)
cap('Figure 1. Full top-level model exported directly from Simulink. Power conversion runs left to right; grouped MSC and GSC/PLL controls sit below the physical network.')
body('The native PMSM block integrates stator d- and q-axis current states and rotor dynamics. Both converters use average-value three-phase voltage-source models. The 5 MW rating follows Wang et al. (2019); other machine, turbine, filter, grid and controller numbers are project engineering assumptions.')

# Page 2 - controls and mechanism.
title('2  Converter controls and weak-grid response',newpage=True)
body('The MSC uses rotor-speed and d/q current loops to control generator torque. The GSC regulates DC-link voltage and grid current in a PLL-aligned frame. Figure 2 shows the actual control subsystems; measurement and Scope paths are separate from the power path.')
two_images([ASSET/'02_msc_control.png',ASSET/'03_gsc_pll_control.png'],[4.65,3.8])
cap('Figure 2. Machine-side control (left) and GSC/PLL control (right), exported from the saved Simulink model.')
body('A 12 to 13 m/s wind step at 2 s changes turbine torque and converter power. With nominal external short-circuit ratio (SCR) reduced from 10 to 5, the same disturbance excites a larger, decaying PCC q-axis voltage oscillation. The oscillation is a coupled control/network response; no sinusoid is injected. Its weak-grid FFT peak is 4.681 Hz, and the near-equilibrium physical-model snapshot mode is 4.638 Hz with 5.24% damping.')
body('The tested mitigation changes only PLL proportional gain from 0.00558 to 0.02 rad/(s V). The integral gain remains 1.56 rad/(s² V). Wang et al., DOI 10.1109/EI247390.2019.9061793, motivates Prony-based SSO analysis; these measured numbers are from this model.')

# Page 3 - results.
title('3  Results and interpretation',newpage=True)
body('Figure 3 compares identical wind steps with common voltage and time axes. Figure 4 is a captured native Simulink Scope display from the weak-grid run. The high-resolution comparison and the Scope both come from the saved physical model.')
table=doc.add_table(rows=1,cols=2);table.autofit=False
table.cell(0,0).width=Inches(7.05);table.cell(0,1).width=Inches(3.3)
for cell in table.rows[0].cells:cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
table.cell(0,0).paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
table.cell(0,0).paragraphs[0].add_run().add_picture(str(sized(ASSET/'08_scenario_comparison.png',2200)),width=Inches(6.95))
scope_path=SCOPE/'weak_grid_11.png'
assert scope_path.is_file(),scope_path
table.cell(0,1).paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
table.cell(0,1).paragraphs[0].add_run().add_picture(str(sized(scope_path,1200)),width=Inches(3.1))
cap('Figure 3. MATLAB comparison of PCC q-axis voltage. Figure 4. Native Simulink Scope snapshot from the weak-grid response.')
t=doc.add_table(rows=2,cols=4);t.style='Light Shading Accent 1'
for i,x in enumerate(('Case','Early PCC q std','Late DC voltage','Late PCC power')):t.cell(0,i).text=x
for i,x in enumerate(('SCR 10 / SCR 5 / mitigated','1.788 / 3.997 / 0.554 V','1405.27 / 1405.42 / 1405.42 V','4.916 / 4.916 / 4.916 MW')):t.cell(1,i).text=x
body('The early standard deviation uses 2.25–3.25 s. Relative to the unmitigated weak-grid run, PLL retuning reduces that metric by a factor of 7.21. Damping of the matched approximately 4.6 Hz linearized mode rises from 5.24% to 18.99%. The mitigated least-damped mode overall is a different 1.843 Hz mode. All four scenario runs and 25 physical/numerical checks pass.')
body('These are responses of an assumed, averaged single-turbine system. The model does not establish a sustained instability, a field-calibrated SCR threshold, PWM harmonics or shaft torsion. Full methods and parameter provenance are in the accompanying technical and validation reports.')

OUT.mkdir(exist_ok=True)
doc.save(DEST)
print(DEST)

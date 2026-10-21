"""Render current physical-model Markdown reports into editable Word files.

PDFs are emitted from the verified DOCX rendering by the documents renderer.
"""
from pathlib import Path
import subprocess
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'report'
SOURCES={
 'TECHNICAL_REPORT':'Vataleni_PMSG_SSO_Technical_Report',
 'MODEL_GUIDE':'Vataleni_PMSG_SSO_Model_Guide',
 'VERIFICATION_REPORT':'Vataleni_PMSG_SSO_Validation_Report',
}

for source,stem in SOURCES.items():
    md=ROOT/'docs'/(source+'.md');docx=OUT/(stem+'.docx')
    subprocess.run(['pandoc',str(md),'--standalone',
                    '--resource-path='+str(ROOT/'docs'),'-o',str(docx)],check=True)
    doc=Document(docx)
    sec=doc.sections[0]
    sec.top_margin=Inches(.78);sec.bottom_margin=Inches(.72)
    sec.left_margin=Inches(.78);sec.right_margin=Inches(.78)
    styles=doc.styles
    normal=styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(6)
    for name,size in [('Title',20),('Heading 1',14),('Heading 2',11.5)]:
        st=styles[name];st.font.name='Calibri';st.font.size=Pt(size)
        st.font.color.rgb=RGBColor(0,0,0)
        st.font.underline=False
        st.paragraph_format.space_before=Pt(12 if name!='Title' else 0)
        st.paragraph_format.space_after=Pt(7)
    for p in doc.paragraphs:
        if p.style.name=='Heading 1':
            p.style=styles['Title'];p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            break
    for table in doc.tables:
        try: table.style='Light Shading Accent 1'
        except KeyError: pass
        table.autofit=True
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(2)
                    for run in p.runs: run.font.size=Pt(8.4)
    for inline in doc.inline_shapes:
        if inline.width>Inches(6.8):
            scale=Inches(6.8)/inline.width
            inline.width=Inches(6.8);inline.height=int(inline.height*scale)
    doc.save(docx)
    print(docx)

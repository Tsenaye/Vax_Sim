from __future__ import annotations

import csv
from pathlib import Path
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "report"
FIG = ROOT / "figures"
SNAP = ROOT / "artifacts" / "model_snapshots"
RESULTS = ROOT / "results"
NAVY = RGBColor(31, 77, 120)
BLUE = RGBColor(46, 116, 181)
MUTED = RGBColor(89, 89, 89)
LIGHT = "E8EEF5"


def read_rows(path: Path):
    with path.open(newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


SUMMARY = read_rows(RESULTS / "summary_cases.csv")
MIT = read_rows(RESULTS / "mitigation_comparison.csv")
MODAL = read_rows(RESULTS / "modal_comparison.csv")
BEFORE = SUMMARY[0]
AFTER = SUMMARY[1]


def set_font(run, name="Calibri", size=11, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def configure(doc: Document, title: str):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"; normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8), ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, NAVY, 8, 4)]:
        st = doc.styles[name]; st.font.name = "Calibri"; st.font.size = Pt(size)
        st.font.color.rgb = color; st.font.bold = True
        st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after)
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(header.add_run(title), size=9, color=MUTED)
    add_page_field(sec.footer.paragraphs[0])


def cover(doc: Document, kicker: str, title: str, subtitle: str):
    for _ in range(5): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18); set_font(p.add_run(kicker.upper()), size=10, color=BLUE, bold=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8); set_font(p.add_run(title), size=28, color=NAVY, bold=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(40); set_font(p.add_run(subtitle), size=14, color=MUTED)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("Authoritative model: model/Vataleni_PMSG_SSO_Final.slx"), size=10.5, color=NAVY, bold=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(f"Generated from fresh MATLAB/Simulink results - {datetime.now():%d %B %Y}"), size=10, color=MUTED)
    doc.add_page_break()


def para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        set_font(p.add_run(bold_lead), bold=True)
        set_font(p.add_run(text[len(bold_lead):]))
    else:
        set_font(p.add_run(text))
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5); p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(8); p.paragraph_format.line_spacing = 1.167
    set_font(p.add_run(text)); return p


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    hdr = t.rows[0]; set_repeat_header(hdr)
    for i, h in enumerate(headers):
        c = hdr.cells[i]; c.text = ""; c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        c._tc.get_or_add_tcPr().append(shading(LIGHT)); set_cell_margins(c)
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(str(h)), size=9.5, bold=True, color=NAVY)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""; cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cells[i])
            p = cells[i].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(str(value)), size=9.3)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths): row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def shading(fill):
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); return shd


def equation(doc, text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(8)
    set_font(p.add_run(text), name="Cambria Math", size=10.5, color=NAVY, italic=True)


def figure(doc, filename, caption, width=5.5):
    path = Path(filename)
    if not path.is_absolute(): path = ROOT / filename
    if path.is_file():
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(); shape = r.add_picture(str(path), width=Inches(width))
        shape._inline.docPr.set('descr', caption)
        shape._inline.docPr.set('title', path.stem.replace('_', ' '))
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_together = True
        set_font(p.add_run(caption), size=9, color=MUTED, italic=True)


def technical_report():
    doc = Document(); configure(doc, "Vataleni PMSG + SSO Technical Report")
    cover(doc, "Technical report", "Vataleni PMSG + SSO Wind-Energy System",
          "Self-contained Simulink implementation, oscillation study, mitigation and validation")
    doc.add_heading("Executive summary", 1)
    para(doc, "This report documents the repaired 5 MW averaged dq wind-energy model. The authoritative SLX contains the wind-turbine, PMSG rotor, machine-side control, DC link, grid-side control, PLL, grid-current dynamics, PCC algebra and measurement equations internally. It does not call custom Level-2 MATLAB S-Functions and it runs from a fresh MATLAB session without an initialization script.")
    para(doc, f"The weak-grid disturbance changes short-circuit ratio (SCR) from 10 to 5 at 2 s. The fresh simulation identifies a dominant electrical-power oscillation at {float(BEFORE['FFT_Hz']):.4f} Hz. The matching nonlinear-state Jacobian mode is {float(MODAL[0]['EigenReal_per_s']):+.5f} +/- j{2*3.141592653589793*float(MODAL[0]['EigenFrequency_Hz']):.5f} per second, with negative damping ratio {float(BEFORE['EigenDampingRatio']):.5f}. Reducing the PLL integral gain from 2.8 to 0.5 moves the tracked mode into the stable half-plane and reduces late-window power-oscillation RMS from {float(BEFORE['PowerOscRMS_pu']):.6f} pu to {float(AFTER['PowerOscRMS_pu']):.3e} pu.")
    table(doc,["Item","Validated value"],[
        ["MATLAB / Simulink","R2025b / 25.2"],["Model","model/Vataleni_PMSG_SSO_Final.slx"],
        ["Solver","ode4, fixed step, 0.0002 s"],["Simulation horizon","12 s"],
        ["Custom S-Functions","0"],["Algebraic loops","0"]],[2.2,4.3])

    doc.add_heading("1. Objective and study scope",1)
    para(doc,"The model studies converter-grid interaction in a full-converter permanent-magnet synchronous-generator wind turbine. It is intended for time-domain and small-signal investigation of sub-synchronous oscillation (SSO) under changing grid strength, and for a controlled comparison of a PLL-gain mitigation.")
    bullet(doc,"Base rating: 5 MVA, 690 V line-line, 50 Hz.")
    bullet(doc,"Plant representation: nonlinear averaged dq model; no switching ripple.")
    bullet(doc,"Disturbance: SCR step from 10 to 5 at 2 s.")
    bullet(doc,"Mitigation: PLL integral gain changes from 2.8 to 0.5 with all other case parameters held equal.")

    doc.add_heading("2. System architecture",1)
    para(doc,"The top-level diagram is organized left-to-right. Compact diagnostic vectors feed the logging interface without altering the physical signal paths. Seven equation groups use connected native Simulink math blocks and every continuous state is an Integrator. One embedded MATLAB Function block assembles the diagnostic vector and reconstructed phase waveforms.")
    figure(doc,SNAP/"top_level.png","Figure 1. Authoritative top-level Simulink architecture.",6.4)
    table(doc,["Subsystem","Inputs","Outputs / states"],[
        ["Wind Turbine + PMSG","wind speed, machine iq","rotor speed, mechanical power/torque; omega_r"],
        ["MSC Control","rotor speed","machine iq and reference; iq_m, speed PI state"],
        ["DC Link","electromagnetic torque, speed, grid power","vdc; capacitor-energy state"],
        ["GSC Control","vdc, id, iq, generator power","vcd/vcq commands; three PI states"],
        ["PLL","PCC vq, mitigation selector","angle, frequency deviation; angle and PI states"],
        ["Grid Current","converter voltage, PLL frame, SCR","id/iq; two current states"],
        ["PCC Interface","angle, id, iq, SCR","vd/vq, P/Q, voltage magnitude"],
        ["Measurements","diagnostic vectors and time","35 logged engineering channels"]],[1.35,2.25,2.9])

    doc.add_heading("3. Governing equations",1)
    doc.add_heading("3.1 Aerodynamics and rotor",2)
    para(doc,"The standard exponential power-coefficient fit is bounded to 0 <= Cp <= 0.50. Aerodynamic power is limited to 1.15 pu and torque follows from power divided by pu speed.")
    equation(doc,"lambda = omega_r omega_base R / v_w ;  P_m = 0.5 rho A Cp v_w^3 / S_base")
    equation(doc,"d omega_r / dt = [T_m - i_qm - D(omega_r - 1)] / (2H)")
    doc.add_heading("3.2 Machine-side converter",2)
    equation(doc,"i_qm,ref = Kp_speed(omega_r - 1) + xi_speed")
    equation(doc,"d i_qm/dt = (i_qm,ref - i_qm)/tau_i ;  d xi_speed/dt = Ki_speed(omega_r - 1)")
    para(doc,"The machine electrical inner loop is an averaged first-order torque-current channel. Machine d-axis current is zero in this reduced-order representation; switching-level machine voltage references are therefore outside the model fidelity.")
    doc.add_heading("3.3 DC link",2)
    equation(doc,"d v_dc/dt = (T_e omega_r - P_grid) / [C_pu max(v_dc, 0.25)]")
    para(doc,"Generator electrical input is electromagnetic torque times rotor speed. The reduced-order machine assumes ideal conversion between air-gap power and the DC link; explicit stator copper and converter losses are outside this model.")
    doc.add_heading("3.4 Grid-side converter",2)
    equation(doc,"i_d,ref = P_gen + Kp_dc(v_dc - v_dc,ref) + xi_dc ;  i_q,ref = 0")
    equation(doc,"v_cd = E + Kp_i(i_d,ref - i_d) + xi_id - X_f i_q")
    equation(doc,"v_cq = Kp_i(i_q,ref - i_q) + xi_iq + X_f i_d")
    doc.add_heading("3.5 PLL and weak-grid interface",2)
    equation(doc,"Delta omega_PLL = Kp_PLL v_q + xi_PLL ;  d delta/dt = omega_base Delta omega_PLL")
    equation(doc,"d xi_PLL/dt = Ki_PLL v_q ;  Ki_PLL = 2.8 (base) or 0.5 (mitigated)")
    equation(doc,"v_d = E cos(delta) + R_g i_d - X_g i_q ;  v_q = -E sin(delta) + R_g i_q + X_g i_d")
    equation(doc,"P = v_d i_d + v_q i_q ;  Q = v_q i_d - v_d i_q")
    doc.add_heading("3.6 Grid current dynamics",2)
    equation(doc,"d i_d/dt = omega_b/L_t [v_cd - E cos(delta) - R_t i_d + L_t(1+Delta omega_PLL)i_q]")
    equation(doc,"d i_q/dt = omega_b/L_t [v_cq + E sin(delta) - R_t i_q - L_t(1+Delta omega_PLL)i_d]")

    doc.add_heading("4. Parameterization and initialization",1)
    table(doc,["Group","Parameter","Value"],[
        ["Base","S, Vll, f","5 MVA; 690 V; 50 Hz"],["Wind","rho, R, lambda_opt","1.225 kg/m3; 55 m; 6.3"],
        ["Rotor","H, D","4.0 s; 0.02 pu"],["MSC","Kp/Ki speed; tau_i","2.0 / 1.0; 0.02 s"],
        ["DC","C_pu; vdc_ref","0.08 s; 1.0 pu"],["GSC","Kp_i/Ki_i; Kp_dc/Ki_dc","0.30/20; 0.80/20"],
        ["Filter","R_f; X_f","0.01 pu; 0.15 pu"],["Grid","E; X/R","1.0 pu; 10"],
        ["PLL","Kp; Ki base/mitigated","0.01; 2.8 / 0.5"]],[1.0,2.5,3.0])
    para(doc,"Initial conditions are the consistent SCR=10 power-flow solution for nominal wind and 1 pu rotor speed. Those numerical state values, scenario defaults, stop time and fixed step are saved inside the model. No callback, addpath command or base-workspace parameter structure is needed for standalone execution.")

    doc.add_heading("5. Numerical configuration",1)
    para(doc,"The final configuration uses fixed-step classical fourth-order Runge-Kutta (ode4), dt = 0.0002 s, and a 12 s horizon. The model contains eleven continuous states and no zero-crossing-dependent logic. The fixed step resolves the approximately 4.84 Hz SSO with more than one thousand points per cycle while preserving deterministic batch reproducibility.")
    figure(doc,REPORT/"evidence"/"figures"/"07_dt_sensitivity.png","Figure 2. Timestep sensitivity of the FFT-identified SSO frequency.",5.8)
    para(doc,"A fresh sweep over 1.0 ms, 0.5 ms, 0.2 ms and 0.1 ms gives 4.8370, 4.8370, 4.8447 and 4.8447 Hz respectively. The 0.2 ms result is converged to the 0.1 ms rerun at the reported precision.")

    doc.add_heading("6. Scenario definitions",1)
    table(doc,["Scenario","SCR","Wind","PLL Ki","Purpose"],[
        ["Baseline","10 -> 10","12 m/s","2.8","steady strong-grid reference"],
        ["Wind step","10 -> 10","12 -> 10.5 m/s at 4 s","2.8","aerodynamic/control response"],
        ["SSO before","10 -> 5 at 2 s","12 m/s","2.8","weak-grid instability"],
        ["SSO after","10 -> 5 at 2 s","12 m/s","0.5","controlled mitigation comparison"]],[1.15,1.15,1.55,0.8,1.85])

    doc.add_heading("7. Time-domain results",1)
    figure(doc,FIG/"01_baseline_overview.png","Figure 3. Baseline rotor speed, active power and DC-link voltage.")
    figure(doc,FIG/"04_sso_time_domain.png","Figure 4. Weak-grid SSO response in power, DC-link voltage and q-axis grid current.")
    figure(doc,FIG/"15_speed_torque_currents.png","Figure 5. Generator speed, mechanical/electromagnetic torque and grid currents.")
    figure(doc,FIG/"16_pcc_voltage_and_power.png","Figure 6. PCC dq voltage and active/reactive power.")
    figure(doc,FIG/"17_scenario_comparison.png","Figure 7. Baseline, weak-grid SSO and mitigated active-power responses.")

    doc.add_heading("8. Frequency-domain and modal analysis",1)
    para(doc,"FFT analysis uses the simulated active-power channel after the configured analysis start. The nonlinear model is also linearized numerically by a centered finite-difference Jacobian of the exact eleven-state ODE at the case equilibrium. Prony analysis provides an independent damped-sinusoid estimate. No eigenvalues or spectral peaks are loaded from a cache.")
    figure(doc,FIG/"05_sso_fft.png","Figure 8. FFT spectrum of the weak-grid active-power oscillation.")
    figure(doc,FIG/"07_eigenvalue_map.png","Figure 9. Eigenvalue map of the weak-grid state Jacobian.")
    table(doc,["Case","FFT (Hz)","Prony sigma (1/s)","Eigen real (1/s)","Eigen freq. (Hz)","Eigen damping"],[
        ["Before",f"{float(BEFORE['FFT_Hz']):.4f}",f"{float(MODAL[0]['PronySigma_per_s']):+.5f}",f"{float(MODAL[0]['EigenReal_per_s']):+.5f}",f"{float(MODAL[0]['EigenFrequency_Hz']):.4f}",f"{float(MODAL[0]['EigenDampingRatio']):+.5f}"],
        ["After",f"{float(AFTER['FFT_Hz']):.4f}",f"{float(MODAL[1]['PronySigma_per_s']):+.5f}",f"{float(MODAL[1]['EigenReal_per_s']):+.5f}",f"{float(MODAL[1]['EigenFrequency_Hz']):.4f}",f"{float(MODAL[1]['EigenDampingRatio']):+.5f}"]],[0.8,0.95,1.2,1.15,1.15,1.15])

    doc.add_heading("9. Mitigation assessment",1)
    figure(doc,FIG/"11_mitigation_before_after.png","Figure 10. Active-power response before and after PLL-gain mitigation.")
    figure(doc,FIG/"12_mitigation_fft_comparison.png","Figure 11. Spectral comparison before and after mitigation.")
    figure(doc,FIG/"13_mitigation_eigenvalue_comparison.png","Figure 12. Eigenvalue comparison before and after mitigation.")
    table(doc,["Metric","Before","After","Change"],[[r["Metric"],f"{float(r['Before']):.6g}",f"{float(r['After']):.6g}",f"{float(r['Change']):+.6g}"] for r in MIT],[2.5,1.2,1.2,1.2])
    para(doc,"The comparison supports the limited claim that the selected PLL integral-gain reduction stabilizes the tracked converter-grid mode for this parameter set. It is not evidence of universal robustness across turbine ratings, grid topologies, switching models or untested operating points.")

    doc.add_heading("10. Validation and portability",1)
    para(doc,"The validator starts with bdclose('all'), clear and clc. It opens and compiles the authoritative SLX without adding project scripts, executes its default simulation, audits S-Functions, callbacks, unresolved links, machine paths and algebraic loops, then regenerates baseline, SSO, mitigation, wind-step, figures and compact result exports. It throws an error unless all twenty checks pass.")
    bullet(doc,"Core dynamics: seven connected native math subsystems, eleven Integrators, and one internal MATLAB Function block for diagnostic assembly.")
    bullet(doc,"Custom Level-2 MATLAB S-Functions in authoritative model: zero.")
    bullet(doc,"Model callbacks: empty.")
    bullet(doc,"Unresolved library links: zero.")
    bullet(doc,"Algebraic loops: zero; no artificial loop-breaking Memory or Unit Delay blocks are required.")
    bullet(doc,"Standalone default simulation: load_system followed by sim from a clean MATLAB process.")

    doc.add_heading("11. Limitations",1)
    bullet(doc,"The converter and machine are averaged; semiconductor switching, PWM ripple and detailed magnetic saturation are excluded.")
    bullet(doc,"The rotor is a single-mass inertia; shaft torsional modes are not represented.")
    bullet(doc,"The machine-side electrical loop is reduced to a first-order q-axis torque-current response with id = 0.")
    bullet(doc,"Generator stator copper losses and detailed converter losses are not represented in the DC energy balance.")
    bullet(doc,"Mitigation is evaluated for the documented SCR step and nominal operating point, not a full grid-code envelope.")
    bullet(doc,"The study is a model-based investigation and is not a substitute for field, controller-hardware or EMT validation.")

    doc.add_heading("12. Conclusions",1)
    para(doc,"The repaired SLX is a portable, self-contained implementation of the recovered nonlinear averaged model. The weak-grid case produces an internally consistent 4.84 Hz SSO identified by FFT, Prony and eigenanalysis. The tested PLL-gain mitigation changes the critical modal real part from positive to negative and materially reduces DC-link deviation, peak current and active-power oscillation for the controlled comparison. The model, scripts, figures, result tables and reports are regenerated from the same authoritative equations and scenario definitions.")

    doc.add_heading("13. References and parameter provenance",1)
    para(doc,"The literature supports the architecture and analysis methods. Only the 5 MW machine rating is taken as a numeric model value from a cited paper. The 690 V base, rotor, filter, DC-link, network and controller values are documented engineering assumptions or study tuning; the simulated 4.84 Hz mode is a project result, not a published benchmark target. See docs/SOURCE_AUDIT.md and docs/PARAMETERS.md for equation and value traceability.")
    bullet(doc,"C. Wang, C. Liu, J. Yu and S. Xu, 'Research on Sub-synchronous Oscillation Characteristics between PMSG-based Wind Farms and Weak AC Grids Based on Prony Method,' 2019, DOI 10.1109/EI247390.2019.9061793. GSC structure, SCR and Prony methodology; 5 MW rating.")
    bullet(doc,"X. Zhan, C. Liu and Z. Wang, 'Analysis and control on Sub-synchronous oscillation (SSO) of HVDC transmission for large-scale PMSG-based wind farm integration,' 2019, DOI 10.1049/joe.2018.8811. PMSG dq equations and control-mediated SSO context; HVDC-specific controller is not implemented here.")
    bullet(doc,"P. Kundur, Power System Stability and Control, 1994, ISBN 0-07-035958-X. Modal frequency and damping conventions.")
    bullet(doc,"M. Bongiorno, A. Petersson and E. Agneholm, The Impact of Wind Farms on Subsynchronous Resonance in Power Systems, Elforsk report 11:29, 2011. Complementary time, frequency and eigenvalue analysis context.")

    doc.add_page_break()
    doc.add_heading("Appendix A. Logged signal dictionary",1)
    signals=["wind_speed","omega_r","torque_m","power_m","iq_machine","vdc","id_grid","iq_grid","power_e","reactive_q","vd_pcc","vq_pcc","v_pcc","pll_angle","ia","va","current_pu","scr","cp","lambda","turbine_speed","electrical_rotor_speed","torque_e","id_machine","iq_machine_reference","vcd_reference","vcq_reference","pll_frequency_hz","sso_signal","damping_gain","controller_command","grid_voltage_pu","id_grid_reference","iq_grid_reference","pll_ki_selected"]
    split=18
    signal_rows=[]
    for i in range(split):
        left=[i+1,signals[i]]
        j=i+split
        right=[j+1,signals[j]] if j<len(signals) else ['', '']
        signal_rows.append(left+right)
    table(doc,["#","Signal","#","Signal"],signal_rows,[0.45,2.8,0.45,2.8])
    return doc


def explainer():
    doc=Document(); configure(doc,"Vataleni PMSG + SSO Progress Explainer")
    cover(doc,"Supervisor explainer","Vataleni PMSG + SSO","What the system does, what the scenarios mean, and how the results were calculated")
    doc.add_heading("What this system is",1)
    para(doc,"This is a 5 MW full-converter wind-turbine model. Wind power accelerates a PMSG rotor, the machine-side converter controls generator torque, a DC capacitor transfers energy between converters, and the grid-side converter injects controlled dq current into a grid represented by a variable short-circuit ratio.")
    figure(doc,SNAP/"top_level.png","Figure 1. The repaired model's eight main engineering subsystems.",6.4)
    doc.add_heading("What each subsystem does",1)
    for lead,text in [
        ("Wind Turbine + PMSG: ","converts wind speed into mechanical power and rotor acceleration."),
        ("MSC Control: ","uses rotor-speed error to command electromagnetic torque current."),
        ("DC Link: ","integrates the energy difference between machine input and grid output."),
        ("GSC Control: ","regulates DC voltage and dq grid current."),
        ("PLL: ","estimates grid angle and frequency from PCC q-axis voltage."),
        ("Grid Current: ","integrates filter and weak-grid dq current equations."),
        ("PCC Interface: ","computes terminal voltage and active/reactive power."),
        ("Measurements: ","exports 35 named channels used by the plots and metrics.")]:
        para(doc,lead+text,bold_lead=lead)
    doc.add_page_break()
    doc.add_heading("The three main scenarios",1)
    table(doc,["Scenario","Meaning","Expected reading"],[
        ["Baseline","Strong grid, SCR 10","stable reference response"],
        ["SSO","Grid weakens from SCR 10 to 5","a growing converter-grid oscillation appears"],
        ["Mitigation","Same weak grid; PLL Ki reduced","tracked mode becomes stable and oscillation decays"]],[1.1,2.8,2.6])
    figure(doc,FIG/"17_scenario_comparison.png","Figure 2. Direct comparison of baseline, SSO and mitigation.",5.55)
    doc.add_heading("What the main graphs show",1)
    figure(doc,FIG/"04_sso_time_domain.png","Figure 3. SSO appears together in active power, DC voltage and q-axis current.",5.55)
    figure(doc,FIG/"12_mitigation_fft_comparison.png","Figure 4. The dominant weak-grid spectral peak is removed by the tested mitigation.",5.55)
    doc.add_heading("How the reported values were calculated",1)
    bullet(doc,f"FFT: active power after the analysis start gives {float(BEFORE['FFT_Hz']):.4f} Hz before mitigation.")
    bullet(doc,f"Prony: a fitted damped sinusoid gives sigma = {float(MODAL[0]['PronySigma_per_s']):+.5f} 1/s before and {float(MODAL[1]['PronySigma_per_s']):+.5f} 1/s after.")
    bullet(doc,f"Eigenanalysis: the numerical Jacobian gives real part {float(MODAL[0]['EigenReal_per_s']):+.5f} 1/s before and {float(MODAL[1]['EigenReal_per_s']):+.5f} 1/s after.")
    bullet(doc,f"Time-domain RMS: the late active-power oscillation changes from {float(BEFORE['PowerOscRMS_pu']):.6f} pu to {float(AFTER['PowerOscRMS_pu']):.3e} pu.")
    bullet(doc,f"DC-link excursion: maximum deviation changes from {float(BEFORE['VdcMaxDeviation_pu']):.4f} pu to {float(AFTER['VdcMaxDeviation_pu']):.4f} pu.")
    doc.add_heading("What was repaired",1)
    para(doc,"The earlier SLX contained Level-2 MATLAB S-Function blocks whose external files were not available when a user opened the model. The seven plant and controller equation groups now use connected native Simulink math blocks and native Integrators. One internal MATLAB Function block assembles diagnostics. The stop time, fixed step, scenario defaults and initial conditions are stored in the model, so pressing Run requires no initialization script or pre-populated base workspace.")
    doc.add_heading("Interpretation",1)
    para(doc,"The agreement of FFT, Prony and eigenanalysis around the pre-mitigation 4.84 Hz mode supports the conclusion that the oscillation is generated by the model dynamics rather than by plotting artifacts. The mitigation result is strong for the tested operating point, but the model is averaged and reduced-order; broader robustness and hardware conclusions require additional studies.")
    return doc


def write_trace():
    fields=["Case","SCR","Mitigation","FFT_Hz","PronySigma_per_s","EigenDampingRatio","PowerOscRMS_pu","PeakCurrent_pu","VdcMaxDeviation_pu"]
    with (REPORT/"results_traceability.csv").open("w",newline="",encoding="utf-8") as f:
        w=csv.DictWriter(f,fieldnames=fields); w.writeheader()
        for row in SUMMARY[:2]: w.writerow({k:row[k] for k in fields})


if __name__ == "__main__":
    REPORT.mkdir(exist_ok=True)
    technical_report().save(REPORT/"Vataleni_PMSG_SSO_Technical_Report.docx")
    explainer().save(REPORT/"Vataleni_PMSG_SSO_Progress_Explainer.docx")
    write_trace()
    print("DOCX reports and traceability table generated.")
